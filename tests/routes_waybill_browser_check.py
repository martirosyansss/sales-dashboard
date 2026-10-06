# -*- coding: utf-8 -*-
"""Проверка «Բեռնագիր» (ответ владельца №57) на странице «Развоз» в настоящем браузере.

Запуск из корня проекта:  python tests/routes_waybill_browser_check.py
Имя без префикса test_: pytest его не собирает (нужны Playwright с Chromium и интернет — Leaflet и Excel-библиотека с CDN).
Приложение — то же, что в tests/routes_dispatch_browser_check.py (настоящие шаблон, статика и blueprint, поддельная ERP,
синтетический день), строки заказов — подделка waybill_loader. Порт 8767 на 127.0.0.1 (8766 — у общей проверки «Развоза»).

A после «Կազմել երթերը» у каждой карточки машины три кнопки «Վարորդ», «Բեռնագիր» и «Excel», и у свёрнутой карточки тоже;
R «Վարորդ» (№62): диалог «Վարորդ և առաքիչ» с подсказкой («с этого дня» или «только этот прошедший день»), фокус в
  списке водителя; списки водителя и առաքիչ — ERP + «+ Նոր …»; ничего не выбрано — закрывается без записи; один человек на
  обе роли — ошибка; «+ Նոր վարորդ» — поле имени; имя (с разметкой — текстом) по Enter сохраняется,
  уведомление, имя в шапке карточки, фокус обратно на кнопку; в накладной — в шапке листа и у подписи, в Excel — строка;
B «Բեռնագիր» открывает окно с листом `.sheet` на каждый рейс машины: «ԲԵՌՆԱԳԻՐ», машина, товары подделки (имя, код,
  «N փաթեթ + M հատ»), итог кг, подписи; разметка из ERP экранирована (имя товара с <b> — текстом);
C «Excel» скачивает bernagir_<машина>_<день>.xlsx: лист на рейс, строка заголовка таблицы и товары; фокус остаётся на кнопке;
R2 у каждого свой срок: подмена водителя на день и առաքիչ «с этого дня» никого — подмена остаётся подменой, назавтра
  прежний водитель, առաքիչ нет; в шапке карточки «(փոխարինող)»;
F тот же номер плана, но на сервере у магазина рейса появился заказ (ERP перечитана) → страница сама видит другой состав
  рейса (basis): ошибка «Թարմացրեք էջը», окно закрыто, сервер ответил 200;
G пока шёл запрос, план на экране сменился (ответ с другим rev) → ошибка, файл не скачан;
D план изменён за спиной страницы (закрепление рейса) → «Բեռնագիր»: окно закрывается, ошибка «Թարմացրեք էջը» с кнопкой
  перечитать; Excel — та же ошибка, файл не скачан;
M 1280 px (две колонки): у кнопок только значки, подписи в aria-label, кнопки внутри карточки;
E телефон 390×860: кнопки строкой под шапкой карточки, горизонтальной прокрутки нет;
ошибки страницы (pageerror) и консоли — провал (кроме сетевых для внешних ресурсов и road-lines и двух ожидаемых 409 в D).
"""
from __future__ import annotations

import dataclasses
import io
from datetime import datetime
import logging
import sys
import tempfile
import threading
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'tests'))

import routes_dispatch_browser_check as base  # noqa: E402  (до Flask: задаёт ROUTES_OSM_PATH и ключ AI)
from openpyxl import load_workbook  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from werkzeug.serving import make_server  # noqa: E402

from route_optimizer import views  # noqa: E402
from route_optimizer import waybill as wb  # noqa: E402
from test_route_optimizer import _dorder  # noqa: E402

PORT = 8767
BASE = f'http://127.0.0.1:{PORT}'
DAY = base.DAY
PRODUCTS = {10: wb.Product(10, '2801', 'Գառնի 6լ <b>x</b>', 'հատ', 6.03, 2),
            11: wb.Product(11, '1113', 'Կոլա 1.5լ', 'հատ', 1.65, 6)}


def loader(isns):
    """Каждый заказ: 10 бутылей 6 л и 13 колы (2 упаковки + 1 шт.); первый заказ — «по накладной»."""
    first = sorted(i.upper() for i in isns)[:1]
    return wb.Lines({i.upper(): ((10, 10.0), (11, 13.0)) for i in isns}, frozenset(first), PRODUCTS)


def main() -> int:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    logging.getLogger('werkzeug').setLevel(logging.ERROR)
    results = []

    def check(cond, msg):
        results.append(bool(cond))
        print(('OK   ' if cond else 'FAIL ') + msg)

    app = base.build_app(tempfile.mkdtemp(prefix='waybill-check-'), base.FakeClient())
    app.extensions['route_optimizer'].waybill_loader = loader
    app.extensions['route_optimizer'].driver_list_loader = lambda since, until: ['Ավակիմյան Արթուր', 'Վարդանյան Գարիկ']
    views._clock = lambda: datetime(2026, 9, 30, 18, 0)     # «сейчас» — накануне DAY: день не прошёл, у водителя есть выбор срока
    server = make_server('127.0.0.1', PORT, app, threaded=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    errors = []
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            ctx = browser.new_context(viewport={'width': 1440, 'height': 950}, accept_downloads=True)
            page = ctx.new_page()
            page.on('pageerror', lambda e: errors.append('pageerror: ' + str(e)))
            expected_409 = []      # D: ответы 409 на устаревший план — ожидаемые, браузер пишет их в консоль

            def on_console(m):
                url = (m.location or {}).get('url', '')
                if m.type == 'error' and '409' in m.text and '/api/routes/dispatch/waybill' in url:
                    expected_409.append(url)
                elif m.type == 'error' and not base.is_ignorable(m):
                    errors.append('console: ' + m.text)
            page.on('console', on_console)
            page.on('dialog', lambda d: d.accept())
            for pat in ('https://tiles.api-maps.yandex.ru/**', 'https://*.tile.openstreetmap.org/**'):
                page.route(pat, lambda r: r.fulfill(status=200, content_type='image/png', body=base.TILE_PNG))
            page.goto(f'{BASE}/routes/dispatch?date={DAY}')
            page.wait_for_selector('#dpBody', state='visible', timeout=30000)
            page.wait_for_function("() => !document.getElementById('dpBuild').disabled", timeout=15000)
            page.click('#dpBuild')
            page.wait_for_selector('#dpTruckCards .dp-tcard', timeout=30000)

            # A
            cards = page.locator('#dpTruckCards .dp-tcard')
            n = cards.count()
            ok = n > 0 and all(cards.nth(i).locator('.dp-tacts .dp-wbbtn').count() == 3 for i in range(n))
            check(ok, f'A {n} truck cards, each with 3 buttons (driver, waybill, Excel)')
            card = cards.first
            truck = card.get_attribute('data-truck')
            labels = card.locator('.dp-wbbtn').evaluate_all("els => els.map(b => b.textContent.trim() + '|' + b.getAttribute('aria-label'))")
            check(labels[0].startswith('Վարորդ|Վարորդ և առաքիչ՝') and labels[1].startswith('Բեռնագիր|Տպել բեռնագիրը՝')
                  and labels[2].startswith('Excel|Բեռնագիրը Excel-ով՝') and truck in labels[1], f'A labels {labels}')
            if card.locator('.dp-thead').get_attribute('aria-expanded') == 'true':
                card.locator('.dp-thead').click()
            check(card.locator('.dp-thead').get_attribute('aria-expanded') == 'false'
                  and card.locator('.dp-wbbtn').nth(1).is_visible(), 'A buttons visible on a collapsed card')

            # R
            DRIVER = 'Վարդանյան Գարիկ <i>x</i>'
            HELPER = 'Ավակիմյան Արթուր'
            card.locator('.dp-drvbtn').click()
            page.wait_for_selector('#dpDriverDlg[open]', timeout=5000)
            past = page.request.get(f'{BASE}/api/routes/dispatch', params={'date': DAY}).json()['is_past']
            check(('Միայն' if past else 'Նախորդ օրերի') in page.inner_text('#dpDriverHint')
                  and page.evaluate("() => document.activeElement && document.activeElement.id") == 'dpDriverPick',
                  f'R dialog: hint ({"only that past day" if past else "from this day"}), focus in the driver list')
            boxes = [page.locator(f'#{b}').is_visible() for b in ('dpDriverOnlyDay', 'dpHelperOnlyDay')]
            check(boxes == [not past, not past] and (past or not (page.is_checked('#dpDriverOnlyDay') or page.is_checked('#dpHelperOnlyDay'))),
                  'R per-role «only this day» boxes: hidden on a past day, unchecked when nobody yet')
            want = ['', 'Ավակիմյան Արթուր', 'Վարդանյան Գարիկ', '__new__']
            opts = page.locator('#dpDriverPick option').evaluate_all("els => els.map(o => o.value)")
            hopts = page.locator('#dpHelperPick option').evaluate_all("els => els.map(o => o.value)")
            groups = page.locator('#dpDriverPick optgroup').evaluate_all("els => els.map(g => g.label)")
            check(opts == want and hopts == want and groups == ['ERP-ի առաքիչներ'] and not page.is_visible('#dpDriverNewBox'),
                  f'R driver and helper lists from ERP + «new»: {opts} {hopts} {groups}')
            page.click('#dpDriverSave')                                    # ничего не выбрано — нечего сохранять
            page.wait_for_function("() => !document.getElementById('dpDriverDlg').open", timeout=5000)
            none = page.request.get(f'{BASE}/api/routes/dispatch', params={'date': DAY}).json()
            check(none['drivers'] == {} and none['helpers'] == {}, 'R nothing chosen: dialog closes, nothing saved')
            card.locator('.dp-drvbtn').click()
            page.wait_for_selector('#dpDriverDlg[open]', timeout=5000)
            page.select_option('#dpDriverPick', '__new__')
            check(page.is_visible('#dpDriverNewBox') and not page.is_visible('#dpHelperNewBox')
                  and page.evaluate("() => document.activeElement.id") == 'dpDriverPick',
                  'R «+ Նոր վարորդ» shows the name field (focus stays in the list)')
            page.fill('#dpDriverName', 'Ավակիմյան Արթուր')
            page.select_option('#dpHelperPick', HELPER)
            page.click('#dpDriverSave')                                    # один человек на обе роли — нельзя
            check(page.inner_text('#dpDriverErr') == 'Վարորդն ու առաքիչը նույն մարդն են'
                  and page.get_attribute('#dpHelperPick', 'aria-invalid') == 'true', 'R same person for both: error, nothing sent')
            page.fill('#dpDriverName', '  Վարդանյան   Գարիկ <i>x</i> ')
            page.press('#dpDriverName', 'Enter')
            page.wait_for_function("() => !document.getElementById('dpDriverDlg').open", timeout=10000)
            page.wait_for_function("() => /վարորդ՝/.test((document.getElementById('dpToast') || {}).textContent || '')", timeout=10000)
            head = card.locator('.dp-tdriver').inner_text()
            check(DRIVER in head and HELPER in head and card.locator('.dp-tdriver i:text-is("x")').count() == 0,
                  f'R card header shows driver and helper as text: {head!r}')
            check(page.evaluate("() => document.activeElement && document.activeElement.classList.contains('dp-drvbtn')"),
                  'R focus back on «Վարորդ»')
            card.locator('.dp-drvbtn').click()           # в машине уже есть люди — по умолчанию «только этот день» (подмена)
            page.wait_for_selector('#dpDriverDlg[open]', timeout=5000)
            groups = page.locator('#dpDriverPick optgroup').evaluate_all("els => els.map(g => g.label)")
            check(page.is_checked('#dpDriverOnlyDay') and page.is_checked('#dpHelperOnlyDay') and page.input_value('#dpDriverPick') == DRIVER
                  and page.input_value('#dpHelperPick') == HELPER and groups == ['ERP-ի առաքիչներ', 'Ավելացված ծրագրում'],
                  'R reopen: substitute preselected, current driver (own group) and helper selected')
            page.click('#dpDriverCancel')
            page.wait_for_function("() => !document.getElementById('dpDriverDlg').open", timeout=5000)
            saved = page.request.get(f'{BASE}/api/routes/dispatch', params={'date': DAY}).json()
            check(saved['drivers'].get(truck) == DRIVER and saved['helpers'].get(truck) == HELPER
                  and saved['driver_list'][-1] == {'name': DRIVER, 'erp': False}, f'R saved on the server {saved["drivers"]} {saved["helpers"]}')

            # B
            with ctx.expect_page(timeout=15000) as popup:
                card.locator('.dp-wbbtn').nth(1).click()
            sheet = popup.value
            sheet.wait_for_function("() => document.querySelectorAll('.sheet').length > 0", timeout=15000)
            api = page.request.get(f'{BASE}/api/routes/dispatch/waybill', params={'date': DAY, 'truck': truck}).json()
            n_trips = len(api['trips'])
            text = sheet.inner_text('body')
            check(sheet.locator('.sheet').count() == n_trips and n_trips >= 1, f'B print window: {sheet.locator(".sheet").count()} sheets for {n_trips} trips')
            check('ԲԵՌՆԱԳԻՐ' in text and truck in text and 'Բաց թողեց (պահեստապետ)' in text and 'Ընդունեց (վարորդ)' in text,
                  'B title, truck, signatures')
            check('Գառնի 6լ <b>x</b>' in text and sheet.locator('.sheet b:text-is("x")').count() == 0, 'B ERP markup shown as text')
            check(f'Վարորդ՝ {DRIVER}' in text and f'Ընդունեց (վարորդ)՝ {DRIVER}' in text
                  and sheet.locator('.sheet i:text-is("x")').count() == 0, 'B driver in the sheet header and at the signature')
            check(f'Առաքիչ՝ {HELPER}' in text and f'Ընդունեց (առաքիչ)՝ {HELPER}' in text
                  and sheet.locator('.sheet .sign > div').count() == 3, 'B helper in the header, third signature')
            cola = sheet.locator('.sheet').first.locator('tr', has_text='Կոլա 1.5լ').inner_text().replace(' ', ' ')
            first = api['trips'][0]
            q = next(r for r in first['rows'] if r['product_id'] == 11)
            check(f'{q["packs"]} փաթեթ' in cola and (q['loose'] == 0 or f'+ {q["loose"]} հատ' in cola) and '6-ական' in cola,
                  f'B cola row packs: {cola!r}')
            check(first['invoiced'] >= 0 and 'Քանակները՝' in text, 'B source note present')
            sheet.close()

            # C
            with page.expect_download(timeout=15000) as dl:
                card.locator('.dp-wbbtn').nth(2).click()
            d = dl.value
            data = Path(d.path()).read_bytes()
            book = load_workbook(io.BytesIO(data))
            safe = ''.join(ch for ch in truck if ch.isascii() and ch.isalnum())
            check(d.suggested_filename == f'bernagir_{safe}_{DAY}.xlsx', f'C file name {d.suggested_filename}')
            check(book.sheetnames == [f'Երթ {i + 1}' for i in range(n_trips)], f'C sheets {book.sheetnames}')
            vals = [[c for c in row] for row in book.worksheets[0].iter_rows(values_only=True)]
            head = next((i for i, r in enumerate(vals) if r[0] == '№'), None)
            check(['Վարորդ', DRIVER] in [list(r[:2]) for r in vals] and ['Առաքիչ', HELPER] in [list(r[:2]) for r in vals],
                  'C driver and helper rows in the Excel sheet')
            check(head is not None and vals[head][2] == 'Ապրանք' and {vals[head + 1][2], vals[head + 2][2]} == {'Գառնի 6լ <b>x</b>', 'Կոլա 1.5լ'},
                  'C table header and products in sheet 1')
            check(page.evaluate("() => document.activeElement && document.activeElement.classList.contains('dp-wbbtn')"
                                " && document.activeElement.textContent.trim() === 'Excel'"), 'C focus stays on the Excel button')

            # R2 (ревью: у каждого свой срок) — подмена водителя на день, затем առաքիչ «с этого дня» никого: подмена водителя
            # остаётся подменой, назавтра снова прежний водитель, առաքիչ с этого дня нет
            nxt = '2026-10-02'
            card.locator('.dp-drvbtn').click()
            page.wait_for_selector('#dpDriverDlg[open]', timeout=5000)
            page.select_option('#dpDriverPick', '__new__')
            page.fill('#dpDriverName', 'Սամվել Փոխարինող')
            page.click('#dpDriverSave')
            page.wait_for_function("() => !document.getElementById('dpDriverDlg').open", timeout=10000)
            card.locator('.dp-drvbtn').click()
            page.wait_for_selector('#dpDriverDlg[open]', timeout=5000)
            page.select_option('#dpHelperPick', '')
            page.uncheck('#dpHelperOnlyDay')
            page.click('#dpDriverSave')
            page.wait_for_function("() => !document.getElementById('dpDriverDlg').open", timeout=10000)
            d1 = page.request.get(f'{BASE}/api/routes/dispatch', params={'date': DAY}).json()
            d2 = page.request.get(f'{BASE}/api/routes/dispatch', params={'date': nxt}).json()
            check(d1['drivers'].get(truck) == 'Սամվել Փոխարինող' and truck in d1['substitutes'] and truck not in d1['helpers']
                  and d2['drivers'].get(truck) == DRIVER and truck not in d2['helpers'],
                  f'R2 per-role scope: driver substitute kept, helper none from this day ({d1["drivers"]}, {d2["drivers"]})')
            card.locator('.dp-tdriver').wait_for(timeout=5000)
            check('(փոխարինող)' in card.locator('.dp-tdriver').inner_text(), 'R2 card header marks the substitute')

            # F
            state = app.extensions['route_optimizer']
            loader0 = state.dispatch_loader
            cid = api['trips'][0]['basis'][0][0]
            state.dispatch_loader = lambda since, until, day: dataclasses.replace(
                loader0(since, until, day), orders=loader0(since, until, day).orders + (_dorder(77, cid, 5.0),))
            state.dispatch_cache.clear()
            served = []
            page.on('response', lambda r: served.append(r.status) if '/api/routes/dispatch/waybill' in r.url else None)
            n_pages = len(ctx.pages)
            card.locator('.dp-wbbtn').nth(1).click()
            page.wait_for_selector('#dpActionError:not(.d-none)', timeout=15000)
            page.wait_for_timeout(300)
            check('Թարմացրեք էջը' in page.inner_text('#dpActionErrorText') and served == [200] and len(ctx.pages) == n_pages,
                  f'F same rev, other trip content: page refuses (server {served}), window closed')
            state.dispatch_loader = loader0
            state.dispatch_cache.clear()
            page.evaluate("() => document.getElementById('dpActionError').classList.add('d-none')")

            # G
            def newer_rev(route):
                resp = route.fetch()
                body = resp.json()
                body['rev'] = body['rev'] + 1
                route.fulfill(response=resp, json=body)
            page.route('**/api/routes/dispatch/waybill*', newer_rev)
            got_g = []
            page.on('download', lambda x: got_g.append(x))
            card.locator('.dp-wbbtn').nth(2).click()
            page.wait_for_selector('#dpActionError:not(.d-none)', timeout=15000)
            page.wait_for_timeout(500)
            check(not got_g and 'Թարմացրեք էջը' in page.inner_text('#dpActionErrorText'), 'G plan changed during the request: error, no download')
            page.unroute('**/api/routes/dispatch/waybill*')
            page.evaluate("() => document.getElementById('dpActionError').classList.add('d-none')")

            # D
            day = page.request.get(f'{BASE}/api/routes/dispatch', params={'date': DAY}).json()
            t0 = next(t for t in day['plan']['trucks'] if t['car_code'] == truck)
            pin = page.request.post(f'{BASE}/api/routes/dispatch/edit', data={
                'date': DAY, 'rev': day['rev'], 'action': 'pin', 'trip': t0['trips'][0]['id'], 'truck': truck})
            check(pin.status == 200, f'D plan changed behind the page (pin {pin.status})')
            n_pages = len(ctx.pages)
            card.locator('.dp-wbbtn').nth(1).click()
            page.wait_for_selector('#dpActionError:not(.d-none)', timeout=15000)
            page.wait_for_timeout(300)
            check('Թարմացրեք էջը' in page.inner_text('#dpActionErrorText') and page.locator('#dpActionReload').is_visible()
                  and len(ctx.pages) == n_pages, f'D stale print: error {page.inner_text("#dpActionErrorText")!r}, window closed')
            got = []
            page.on('download', lambda x: got.append(x))
            page.evaluate("() => document.getElementById('dpActionError').classList.add('d-none')")
            card.locator('.dp-wbbtn').nth(2).click()
            page.wait_for_selector('#dpActionError:not(.d-none)', timeout=15000)
            page.wait_for_timeout(500)
            check(not got and 'Թարմացրեք էջը' in page.inner_text('#dpActionErrorText'), 'D stale Excel: error, no download')
            check(len(expected_409) == 2, f'D two 409 answers from the waybill API ({len(expected_409)})')
            page.click('#dpActionReload')
            page.wait_for_timeout(1500)

            # M
            page.set_viewport_size({'width': 1280, 'height': 900})
            page.wait_for_timeout(500)
            c0 = page.locator('#dpTruckCards .dp-tcard').first
            spans = c0.locator('.dp-wbbtn span').evaluate_all("els => els.map(e => e.getBoundingClientRect().width)")
            cb, ab = c0.bounding_box(), c0.locator('.dp-tacts').bounding_box()
            check(all(w <= 1 for w in spans) and ab['x'] + ab['width'] <= cb['x'] + cb['width'] + 0.5
                  and c0.locator('.dp-wbbtn').nth(1).get_attribute('aria-label').startswith('Տպել բեռնագիրը՝'),
                  f'M 1280: icon-only buttons inside the card (label widths {spans})')

            # E
            page.set_viewport_size({'width': 390, 'height': 860})
            page.wait_for_timeout(500)
            if page.locator('#dpTabs').is_visible():   # №81: на телефоне рейсы — во вкладке «Երթեր»
                page.click('.dp-tab[data-tab="trips"]')
                page.wait_for_timeout(300)
            sw = page.evaluate('() => document.documentElement.scrollWidth')
            c0 = page.locator('#dpTruckCards .dp-tcard').first
            hb, ab = c0.locator('.dp-thead').bounding_box(), c0.locator('.dp-tacts').bounding_box()
            check(sw <= 390 and hb and ab and ab['y'] >= hb['y'] + hb['height'] - 1, f'E phone: scrollWidth={sw}, buttons under header')

            check(not errors, 'no pageerror / console errors' + ('' if not errors else ': ' + ' | '.join(errors[:5])))
            browser.close()
    finally:
        server.shutdown()
    print('ALL OK' if all(results) else 'SOME FAILED')
    return 0 if all(results) else 1


if __name__ == '__main__':
    sys.exit(main())
