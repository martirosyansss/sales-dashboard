# -*- coding: utf-8 -*-
"""Проверка страницы «Տարա» (/courier/tare, ответ владельца №87 п. 9) в настоящем браузере.

Запуск из корня проекта:  python tests/courier_tare_browser_check.py
Имя без префикса test_: pytest его не собирает (нужны Playwright с Chromium и интернет — Excel-библиотека с CDN).
Приложение: настоящие шаблон, статика и blueprint «Առաքիչ» на временной courier.db, синтетические два дня трёх магазинов,
ERP подменена (связи тары, магазины по коду). Порт 8771 на 127.0.0.1.

A таблица магазин × тара: больший долг первым, сводка «Ընդամենը խանութներում», период данных;
X название магазина с разметкой — текстом (нет <img> в таблице);
F фильтр «Մենեջեր» и поиск сужают строки, сброс возвращает;
O карандаш у строки → форма в ячейке (количество, дата) → «Պահել»: начальный остаток и баланс обновились;
S клик по магазину → диалог: баланс, история по дням, форма «Սկզբնական մնացորդ» — новый вид тары записан;
I импорт CSV с ошибками → предпросмотр с ошибками, «Գրանցել» выключена; правильный CSV и Excel (дата-ячейка) → запись;
  код магазина из Excel — текстом, как видно в ячейке («0012» с форматом «0000»);
C «CSV» скачивает tara_<день>.csv с заголовком;
P телефон 390 × 860: горизонтальной прокрутки нет, строки — карточками;
ошибки страницы (pageerror) и консоли — провал (кроме сетевых ошибок внешних ресурсов).
"""
from __future__ import annotations

import io
import logging
import sys
import tempfile
import threading
import uuid
from datetime import timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from flask import Flask  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from werkzeug.serving import make_server  # noqa: E402

import courier  # noqa: E402
from courier import clock, events as ev  # noqa: E402
from courier.erp_day import ContainerLink  # noqa: E402

PORT = 8771
BASE = f'http://127.0.0.1:{PORT}'
BOTTLE, CRATE = 'erp:900', 'erp:901'
XSS = 'Խանութ <img src=x onerror="window.__xss=1">'


class FakeDb:
    connection_string = 'DRIVER={none};'


def _stop(sid, cid, name, agent, lines, tare, seq):
    return {'stop_id': sid, 'seq': seq, 'collect': 'none', 'doc_number': f'D{seq}{cid}',
            'customer': {'id': cid, 'code': f'C{cid}', 'name': name}, 'agent_name': agent, 'amount_due': 0,
            'lines': [{'line_id': lid, 'qty': q, 'price': 0, 'product_id': pid, 'weight_kg': kg} for lid, q, pid, kg in lines],
            'tare_expected': [{'tare_id': t, 'name': {BOTTLE: 'Շիշ 19լ', CRATE: 'Արկղ'}[t], 'qty': q} for t, q in tare.items()]}


def _ev(etype, sid, payload, day, hhmm):
    return {'id': str(uuid.uuid4()), 'type': etype, 'stop_id': sid, 'date': day, 'at': f'{day}T{hhmm}:00+04:00',
            'payload': payload}


def build_app(tmp: str) -> Flask:
    app = Flask(__name__, template_folder=str(ROOT / 'templates'), static_folder=str(ROOT / 'static'))
    app.secret_key = 'check'
    app.add_url_rule('/logout', 'logout', lambda: '', methods=['POST'])
    courier.init_app(app, FakeDb(), db_path=str(Path(tmp) / 'courier.db'))
    st = app.extensions['courier']
    st.tare_links_loader = lambda: ((ContainerLink(1, 900, 1.0, 1.0), ContainerLink(2, 901, 6.0, 1.0)),
                                    {900: 'Շիշ 19լ', 901: 'Արկղ'})
    st.customer_code_loader = lambda codes: {}
    store = st.store
    did = store.save_driver(None, 'Արամ', True, '1234', 'admin')
    today = clock.today()
    d1, d2 = (today - timedelta(days=2)).isoformat(), (today - timedelta(days=1)).isoformat()
    for car, day, stops, evs in (
        ('CAR1', d1, [_stop('S:' + str(uuid.uuid4()).upper(), 10, XSS, 'Մենեջեր Ա', [('a', 30, 1, 570.0)], {BOTTLE: 30}, 1)], None),
        ('CAR2', d2, [_stop('S:' + str(uuid.uuid4()).upper(), 11, 'Խանութ Բ', 'Մենեջեր Բ', [('b1', 6, 1, 114.0), ('b2', 12, 2, 18.0)],
                            {BOTTLE: 6, CRATE: 2}, 1),
                      _stop('S:' + str(uuid.uuid4()).upper(), 12, 'Խանութ Գ', 'Մենեջեր Ա', [('c', 5, 1, 95.0)], {BOTTLE: 5}, 2)], None)):
        store.save_day(day, car, stops, 'v1', day + 'T08:00:00+04:00')
        terminal, _ = store.create_terminal('T-' + car, car, 'admin')
        who = ev.Who(terminal.id, car, did, 'Արամ')
        batch = []
        for i, s in enumerate(stops):
            batch.append(_ev('delivery', s['stop_id'], {'lines': [{'line_id': ln['line_id'], 'qty': ln['qty']} for ln in s['lines']]},
                             day, f'1{i}:00'))
            batch.append(_ev('tare', s['stop_id'], {'items': [{'tare_id': BOTTLE, 'qty': 2 + i}]}, day, f'1{i}:05'))
        r = ev.ingest(store, who, batch).json()
        assert not r['rejected'], r['rejected']
    return app


def main() -> int:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    logging.getLogger('werkzeug').setLevel(logging.ERROR)
    results = []

    def check(cond, msg):
        results.append(bool(cond))
        print(('OK   ' if cond else 'FAIL ') + msg)

    app = build_app(tempfile.mkdtemp(prefix='tare-check-'))
    server = make_server('127.0.0.1', PORT, app, threaded=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    errors = []
    today = clock.today()
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            ctx = browser.new_context(viewport={'width': 1440, 'height': 950}, accept_downloads=True)
            page = ctx.new_page()
            page.on('pageerror', lambda e: errors.append('pageerror: ' + str(e)))

            def on_console(m):
                url = (m.location or {}).get('url', '')
                if m.type == 'error' and not (url and not url.startswith(BASE)) and 'Failed to load resource' not in m.text:
                    errors.append('console: ' + m.text)
            page.on('console', on_console)
            page.goto(f'{BASE}/courier/tare')
            page.wait_for_selector('#ctBody .ct-store', timeout=20000)
            rows = page.locator('#ctBody tr')

            # A
            first = rows.first.inner_text()
            check(rows.count() == 4 and 'C10' in first and '28' in first, f'A 4 rows, biggest debt first: {first!r}')
            kpi = page.inner_text('#ctKpis')
            check('Շիշ 19լ' in kpi and 'Արկղ' in kpi and 'Տերմինալների տվյալներ' in page.inner_text('#ctPeriod'), 'A summary and period')
            # X
            check(page.locator('#ctBody img').count() == 0 and '<img' in first
                  and not page.evaluate('() => window.__xss === 1'), 'X store name with markup shown as text')
            # F
            page.select_option('#ctAgent', 'Մենեջեր Բ')
            check(rows.count() == 2 and all('C11' in t for t in rows.all_inner_texts()), 'F manager filter → store 11 only')
            page.select_option('#ctAgent', '')
            page.fill('#ctQ', 'Խանութ Գ')
            check(rows.count() == 1 and 'C12' in rows.first.inner_text(), 'F search by name')
            page.fill('#ctQ', '')
            check(rows.count() == 4, 'F reset')

            # O
            row10 = page.locator('#ctBody tr', has_text='C10')
            row10.locator('.ct-edit').click()
            form = row10.locator('form.ct-oform')
            check(form.is_visible() and page.evaluate('() => document.activeElement.type') == 'number', 'O inline form, focus in qty')
            form.locator('input[type=number]').fill('50')
            form.locator('input[type=date]').fill(today.isoformat())
            form.locator('button[type=submit]').click()
            page.wait_for_function("() => [...document.querySelectorAll('#ctBody tr')].some(r => r.innerText.includes('C10') && r.querySelector('.ct-open small'))",
                                   timeout=10000)
            t10 = page.locator('#ctBody tr', has_text='C10').inner_text()
            check('50' in t10 and today.strftime('%d.%m.%Y') in t10, f'O opening saved and shown: {t10!r}')
            api = page.request.get(f'{BASE}/api/courier/admin/tare').json()
            r10 = next(r for r in api['rows'] if r['customer_id'] == 10)
            check((r10['opening'], r10['balance']) == (50.0, 50.0), f'O balance from opening as of today: {r10["balance"]}')

            # S
            page.locator('#ctBody .ct-store', has_text='Խանութ Բ').first.click()
            page.wait_for_selector('#ctStoreDlg[open] .ct-hist', timeout=10000)
            body = page.inner_text('#ctStoreBody')
            check('Խանութ Բ (C11)' in page.inner_text('#ctStoreTitle') and 'Ըստ օրերի' in body
                  and page.locator('#ctStoreBody .ct-hist tbody tr').count() == 2, 'S store dialog with per-day history')
            page.select_option('#ctAddKind', CRATE)
            page.fill('#ctAddQty', '3')
            page.click('#ctAddForm button[type=submit]')
            page.wait_for_function("() => document.querySelector('#ctStoreBody .ct-olist') && document.querySelector('#ctStoreBody .ct-olist').innerText.includes('Արկղ')",
                                   timeout=10000)
            check('3' in page.inner_text('#ctStoreBody .ct-olist'), 'S opening added from the dialog')
            page.keyboard.press('Escape')
            page.wait_for_function("() => !document.getElementById('ctStoreDlg').open", timeout=5000)

            # I
            bad = 'Խանութի կոդ;Տարա;Քանակ;Ամսաթիվ\r\nC12;Շիշ 19լ;7;01.01.2026\r\nC404;Տակառ;x;2099-01-01\r\n'
            page.set_input_files('#ctFile', files=[{'name': 'bad.csv', 'mimeType': 'text/csv', 'buffer': ('﻿' + bad).encode('utf-8')}])
            page.wait_for_selector('#ctImportDlg[open] .ct-ferr', timeout=10000)
            txt = page.inner_text('#ctImportBody')
            check('Այդ կոդով խանութ չի գտնվել' in txt and 'Տարայի այդպիսի տեսակ չկա' in txt
                  and page.is_disabled('#ctImportApply'), 'I CSV with errors → errors shown, apply disabled')
            page.click('#ctImportCancel')
            good = 'C12;"Շիշ 19լ";7;01.01.2026\r\n'
            page.set_input_files('#ctFile', files=[{'name': 'good.csv', 'mimeType': 'text/csv', 'buffer': good.encode('utf-8')}])
            page.wait_for_function("() => !document.getElementById('ctImportApply').disabled", timeout=10000)
            check('Խանութ Գ' in page.inner_text('#ctImportBody'), 'I preview names the store')
            page.click('#ctImportApply')
            page.wait_for_function("() => !document.getElementById('ctImportDlg').open", timeout=10000)
            page.wait_for_timeout(800)
            api = page.request.get(f'{BASE}/api/courier/admin/tare').json()
            r12 = next(r for r in api['rows'] if r['customer_id'] == 12 and r['tare_id'] == BOTTLE)
            check((r12['opening'], r12['as_of']) == (7.0, '2026-01-01'), f'I CSV applied: {r12}')
            wb = Workbook()
            ws = wb.active
            ws.append(['Կոդ', 'Տարա', 'Քանակ', 'Ամսաթիվ'])
            ws.append(['C11', 'erp:900', 4, today - timedelta(days=3)])
            buf = io.BytesIO()
            wb.save(buf)
            page.set_input_files('#ctFile', files=[{'name': 'x.xlsx', 'buffer': buf.getvalue(),
                                                    'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}])
            page.wait_for_function("() => !document.getElementById('ctImportApply').disabled", timeout=15000)
            page.click('#ctImportApply')
            page.wait_for_function("() => !document.getElementById('ctImportDlg').open", timeout=10000)
            page.wait_for_timeout(800)
            api = page.request.get(f'{BASE}/api/courier/admin/tare').json()
            r11 = next(r for r in api['rows'] if r['customer_id'] == 11 and r['tare_id'] == BOTTLE)
            check((r11['opening'], r11['as_of']) == (4.0, (today - timedelta(days=3)).isoformat()), f'I Excel applied (date cell): {r11}')
            # код-число с форматом «0000» уходит на сервер как его видно в Excel — «0012», без потери нулей
            wb = Workbook()
            ws = wb.active
            ws.append([12, 'erp:900', 1, '01.10.2026'])
            ws['A1'].number_format = '0000'
            buf = io.BytesIO()
            wb.save(buf)
            with page.expect_request(lambda q: q.url.endswith('/api/courier/admin/tare/import') and q.method == 'POST') as req:
                page.set_input_files('#ctFile', files=[{'name': 'z.xlsx', 'buffer': buf.getvalue(),
                                                        'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}])
            sent = req.value.post_data_json['rows']
            page.wait_for_selector('#ctImportDlg[open] .ct-ferr', timeout=10000)
            check([r['code'] for r in sent] == ['0012'], f'I Excel code column taken as text: {sent}')
            page.click('#ctImportCancel')

            # C
            with page.expect_download(timeout=10000) as dl:
                page.click('#ctCsv')
            d = dl.value
            text = Path(d.path()).read_bytes().decode('utf-8')
            check(d.suggested_filename == f'tara_{today.isoformat()}.csv' and text.startswith('﻿Խանութի կոդ;'), f'C CSV {d.suggested_filename}')

            # P
            page.set_viewport_size({'width': 390, 'height': 860})
            page.wait_for_timeout(400)
            sw = page.evaluate('() => document.documentElement.scrollWidth')
            check(sw <= 390 and page.locator('#ctTable thead').is_hidden(), f'P phone 390: scrollWidth={sw}, rows as cards')

            check(not errors, 'no pageerror / console errors' + ('' if not errors else ': ' + ' | '.join(errors[:5])))
            browser.close()
    finally:
        server.shutdown()
    print('ALL OK' if all(results) else 'SOME FAILED')
    return 0 if all(results) else 1


if __name__ == '__main__':
    sys.exit(main())
